"""Document-ingestion pipeline (blueprint Section 5).

capture & quarantine -> preserve original (SHA-256) -> resolve context -> duplicate / revision
detection -> parse by family -> extract clauses & structured assertions -> review queue.
Processing is idempotent per (revision, pipeline_version).
"""
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from fastapi import HTTPException
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app import audit, config
from app.models import (ContractClause, Document, DocumentRevision, ExtractionAssertion, Passage,
                        ProcessingJob, now)


# ------------------------------------------------------------------ 1. capture and quarantine
def validate_upload(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in config.ALLOWED_EXTENSIONS:
        raise HTTPException(415, {"code": "QUARANTINED", "message": f"File type {ext or '(none)'} not accepted"})
    if not data:
        raise HTTPException(422, {"code": "QUARANTINED", "message": "Empty file"})
    if len(data) > config.MAX_UPLOAD_BYTES:
        raise HTTPException(413, {"code": "QUARANTINED", "message": "File too large"})
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise HTTPException(422, {"code": "QUARANTINED", "message": "Content is not a PDF"})
    if ext in (".docx", ".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
        except zipfile.BadZipFile:
            raise HTTPException(422, {"code": "QUARANTINED", "message": "Malformed Office file"})
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise HTTPException(422, {"code": "QUARANTINED", "message": "Macro-enabled content rejected"})
        if (ext == ".docx" and "word/document.xml" not in names) or (ext == ".xlsx" and "xl/workbook.xml" not in names):
            raise HTTPException(422, {"code": "QUARANTINED", "message": "Office package structure invalid"})
    if ext in (".txt", ".md"):
        if b"\x00" in data:
            raise HTTPException(422, {"code": "QUARANTINED", "message": "Binary content in text file"})
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(422, {"code": "QUARANTINED", "message": "Text file is not UTF-8"})
    return ext


# ------------------------------------------------------------------ 2. preserve original
def store_object(data: bytes) -> tuple[str, str]:
    digest = hashlib.sha256(data).hexdigest()
    key = f"{digest[:2]}/{digest}"
    path = config.OBJECT_STORE_DIR / key
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_bytes(data)
    return digest, key


def read_object(key: str) -> bytes:
    return (config.OBJECT_STORE_DIR / key).read_bytes()


# ------------------------------------------------------------------ 4. duplicate / revision detection
@dataclass
class IngestResult:
    outcome: str            # NEW_DOCUMENT | NEW_REVISION | DUPLICATE | REVISION_CONFLICT
    document: Document
    revision: DocumentRevision
    message: str


def register_revision(db: Session, *, actor_id: str, project_id: str, document_number: str, revision_label: str,
                      document_type: str, title: str, filename: str, data: bytes, main_contract_id=None,
                      subcontract_id=None, source_party="EMPLOYER", security_classification="CONFIDENTIAL",
                      issued_at=None, effective_from=None, source_channel="MANUAL_UPLOAD") -> IngestResult:
    validate_upload(filename, data)
    digest, key = store_object(data)

    # Same content anywhere in the project -> duplicate; no new business records (DOC-01).
    dup = db.execute(select(DocumentRevision).join(Document).where(
        Document.project_id == project_id, DocumentRevision.content_hash == digest)).scalar_one_or_none()
    if dup:
        return IngestResult("DUPLICATE", dup.document, dup,
                            f"Identical content already registered as {dup.document.document_number} rev {dup.revision_label}")

    doc = db.scalar(select(Document).where(Document.project_id == project_id, Document.document_number == document_number))
    outcome = "NEW_REVISION"
    if doc is None:
        doc = Document(project_id=project_id, document_number=document_number, document_type=document_type,
                       title=title, main_contract_id=main_contract_id, subcontract_id=subcontract_id,
                       source_party=source_party, security_classification=security_classification)
        db.add(doc)
        db.flush()
        outcome = "NEW_DOCUMENT"

    rev = DocumentRevision(document_id=doc.id, revision_label=revision_label, content_hash=digest, object_key=key,
                           filename=filename, size_bytes=len(data), uploaded_by=actor_id, issued_at=issued_at,
                           effective_from=effective_from, source_channel=source_channel)
    same_label = [r for r in doc.revisions if r.revision_label == revision_label]
    if same_label:
        # Same number + label but different content: never guess which one governs (DOC-04).
        rev.ingestion_state = "CONFLICT"
        rev.conflict_note = (f"Revision label '{revision_label}' already exists with different content "
                             f"(hash {same_label[0].content_hash[:12]}...). Reviewer must resolve.")
        outcome = "REVISION_CONFLICT"
    db.add(rev)
    db.flush()
    audit.record(db, actor_id, f"document.{outcome.lower()}", "DocumentRevision", rev.id, project_id=project_id,
                 after={"document_number": document_number, "revision": revision_label, "sha256": digest,
                        "state": rev.ingestion_state})
    msg = rev.conflict_note if outcome == "REVISION_CONFLICT" else "Registered"
    return IngestResult(outcome, doc, rev, msg)


# ------------------------------------------------------------------ 5. parse by document family
@dataclass
class Block:
    text: str
    page: int | None = None
    cell: str | None = None
    kind: str = "TEXT"


def parse(filename: str, data: bytes) -> list[Block]:
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md"):
        blocks = []
        for p, page_text in enumerate(data.decode("utf-8").split("\f"), start=1):
            blocks += [Block(line, p) for line in page_text.splitlines()]
        return blocks
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        blocks = []
        for p, page in enumerate(reader.pages, start=1):
            blocks += [Block(line, p) for line in (page.extract_text() or "").splitlines()]
        return blocks
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        return [Block(par.text) for par in d.paragraphs] + \
               [Block(" | ".join(c.text for c in row.cells), kind="TABLE_ROW") for t in d.tables for row in t.rows]
    if ext == ".xlsx":
        import openpyxl
        wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        blocks = []
        for ws in wb_f.worksheets:
            wv = wb_v[ws.title]
            for row in ws.iter_rows():
                cells = []
                for c in row:
                    if c.value is None:
                        continue
                    if isinstance(c.value, str) and c.value.startswith("="):
                        val = wv[c.coordinate].value
                        cells.append(f"{c.coordinate}={c.value} [value={val if val is not None else 'not cached'}]")
                    else:
                        cells.append(f"{c.coordinate}={c.value}")
                if cells:
                    first, last = row[0].coordinate, row[-1].coordinate
                    blocks.append(Block("; ".join(cells), cell=f"{ws.title}!{first}:{last}", kind="TABLE_ROW"))
        return blocks
    raise ValueError(ext)


CLAUSE_HEAD = re.compile(r"^\s*(?:Clause\s+)?(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+([A-Z][^.]{2,100})\s*$")


def group_passages(blocks: list[Block]) -> list[dict]:
    """Group lines into passages; a clause heading starts a new passage that absorbs its body."""
    out, cur = [], None

    def flush():
        nonlocal cur
        if cur and cur["text"].strip():
            cur["text"] = cur["text"].strip()
            out.append(cur)
        cur = None

    for b in blocks:
        if b.kind == "TABLE_ROW":
            flush()
            out.append({"text": b.text, "page": b.page, "cell": b.cell, "section": None, "heading": None,
                        "kind": "TABLE_ROW"})
            continue
        line = b.text.rstrip()
        m = CLAUSE_HEAD.match(line)
        if m:
            flush()
            cur = {"text": line, "page": b.page, "cell": None, "section": m.group(1), "heading": m.group(2).strip(),
                   "kind": "CLAUSE"}
        elif not line.strip():
            if cur and cur["kind"] == "TEXT":
                flush()
        else:
            if cur is None:
                cur = {"text": "", "page": b.page, "cell": None, "section": None, "heading": None, "kind": "TEXT"}
            cur["text"] += "\n" + line
    flush()
    return out


# ------------------------------------------------------------------ 6. extraction
CLAUSE_TYPES = [  # ordered: first match wins
    ("DELAY_LD", r"liquidated damages|delay damages"),
    ("NOTICE", r"\bnotice\b|\bnotify\b"),
    ("TESTING", r"pressure test|hydrostatic|hydro-test|\btest|inspection"),
    ("HSE", r"health|safety|\bhse\b|environment|permit to work"),
    ("PAYMENT", r"payment|invoice|retention|certificat"),
    ("VARIATION", r"variation|change order|instruction to change"),
    ("INSURANCE", r"insurance|insure"),
    ("WARRANTY", r"warrant|defects liability|defect"),
    ("INDEMNITY_LIABILITY", r"indemn|liability"),
    ("TERMINATION", r"terminat|suspen"),
    ("DISPUTE", r"dispute|arbitration"),
    ("CONFIDENTIALITY", r"confidential"),
    ("QUALITY", r"quality|itp\b"),
    ("DEFINITIONS", r"definition|interpretation|means"),
]
HIGH_RISK = {"DELAY_LD", "NOTICE", "PAYMENT", "INDEMNITY_LIABILITY", "TERMINATION", "INSURANCE"}
LOW_RISK = {"CONFIDENTIALITY", "DEFINITIONS", "GENERAL"}


def classify_clause(heading: str, text: str) -> str:
    hay = f"{heading}\n{heading}\n{text}".lower()
    # heading carries more weight: try heading alone first
    for kind, pat in CLAUSE_TYPES:
        if re.search(pat, heading.lower()):
            return kind
    for kind, pat in CLAUSE_TYPES:
        if re.search(pat, hay):
            return kind
    return "GENERAL"


def clause_risk(kind: str) -> str:
    return "HIGH" if kind in HIGH_RISK else "LOW" if kind in LOW_RISK else "MEDIUM"


WORDS = {w: i for i, w in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen "
    "seventeen eighteen nineteen twenty".split())}
WORDS.update({"twenty-one": 21, "twenty-eight": 28, "thirty": 30, "forty-five": 45, "sixty": 60, "ninety": 90})
NOTICE_RE = re.compile(r"within\s+(?:(\d+)|([a-z\-]+))(?:\s*\((\d+)\))?\s+(calendar\s+|working\s+|business\s+)?days", re.I)
PCT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:%|per\s*cent|percent)", re.I)
AMT_RE = re.compile(r"(?:INR|Rs\.?|USD|₹)\s?[\d,]+(?:\.\d+)?", re.I)
EXCEPTION_RE = re.compile(r"\b(unless|except|save as|provided that|not\s+be\s+required|shall not)\b", re.I)
INJECTION_RE = re.compile(r"ignore (all |any )?(previous|prior|above) instructions|disregard .{0,30}(polic|instruction)|"
                          r"you are now|system prompt|grant .{0,40}(admin|access|approval)|approve this .{0,20}automatically",
                          re.I)


def extract_assertions(text: str) -> list[dict]:
    found = []
    for m in NOTICE_RE.finditer(text):
        digits = m.group(1) or m.group(3)
        n = int(digits) if digits else WORDS.get((m.group(2) or "").lower())
        if n is None:
            continue
        basis = "WORKING" if (m.group(4) or "").strip().lower() in ("working", "business") else "CALENDAR"
        sentence = _sentence_around(text, m.start())
        conf = 0.92 if digits else 0.78
        if EXCEPTION_RE.search(sentence):
            conf -= 0.25   # exceptions/negations need a human
        found.append({"field_name": "NOTICE_PERIOD_DAYS", "raw_value": sentence, "normalised_value": f"{n}|{basis}",
                      "confidence": round(conf, 3)})
    for m in PCT_RE.finditer(text):
        found.append({"field_name": "PERCENTAGE", "raw_value": _sentence_around(text, m.start()),
                      "normalised_value": m.group(1), "confidence": 0.85})
    for m in AMT_RE.finditer(text):
        found.append({"field_name": "AMOUNT", "raw_value": _sentence_around(text, m.start()),
                      "normalised_value": re.sub(r"[^\d.]", "", m.group(0)), "confidence": 0.8})
    return found


def _sentence_around(text: str, pos: int) -> str:
    start = max(text.rfind(".", 0, pos), text.rfind("\n", 0, pos)) + 1
    end = text.find(".", pos)
    return text[start: end + 1 if end != -1 else len(text)].strip()


# ------------------------------------------------------------------ job runner (idempotent)
def process_revision(db: Session, revision: DocumentRevision, pipeline_version: str = config.PIPELINE_VERSION,
                     fail_after: int | None = None) -> ProcessingJob:
    """Run the pipeline. Re-running a completed job is a no-op; an interrupted job is cleaned and rerun.
    ``fail_after`` exists only to simulate an interruption in tests (DOC-06)."""
    if revision.ingestion_state == "CONFLICT":
        raise HTTPException(409, {"code": "REVISION_CONFLICT", "message": revision.conflict_note})
    job = db.scalar(select(ProcessingJob).where(ProcessingJob.revision_id == revision.id,
                                                ProcessingJob.pipeline_version == pipeline_version))
    if job and job.status == "COMPLETED":
        return job
    if job is None:
        job = ProcessingJob(revision_id=revision.id, pipeline_version=pipeline_version)
        db.add(job)
        db.flush()
    else:   # previous attempt interrupted: remove partial output of THIS job only
        db.execute(delete(ExtractionAssertion).where(ExtractionAssertion.job_id == job.id))
        clause_ids = select(ContractClause.id).join(Passage, ContractClause.source_passage_id == Passage.id).where(
            Passage.job_id == job.id)
        db.execute(delete(ContractClause).where(ContractClause.id.in_(clause_ids)))
        db.execute(delete(Passage).where(Passage.job_id == job.id))
    job.status, job.started_at, job.error_code = "RUNNING", now(), None
    revision.ingestion_state = "PROCESSING"
    db.commit()

    try:
        data = read_object(revision.object_key)
        groups = group_passages(parse(revision.filename, data))
        doc = revision.document
        has_clauses = db.scalar(select(ContractClause.id).where(ContractClause.document_revision_id == revision.id))
        n_clauses = n_assert = n_flags = 0
        for i, g in enumerate(groups):
            if fail_after is not None and i >= fail_after:
                raise RuntimeError("simulated interruption")
            flags = ["EMBEDDED_INSTRUCTION"] if INJECTION_RE.search(g["text"]) else []
            n_flags += bool(flags)
            p = Passage(revision_id=revision.id, job_id=job.id, seq=i, page_number=g["page"],
                        section_reference=g["section"], cell_address=g["cell"], content_type=g["kind"],
                        text=g["text"], flags=flags)
            db.add(p)
            db.flush()
            for a in extract_assertions(g["text"]):
                db.add(ExtractionAssertion(passage_id=p.id, job_id=job.id, method=f"regex-{pipeline_version}", **a))
                n_assert += 1
            # Clauses are extracted once per revision; later pipeline versions never overwrite them (DOC-07).
            if g["kind"] == "CLAUSE" and not has_clauses and doc.document_type in ("GCC", "SPECIAL_CONDITIONS",
                                                                                  "MAIN_CONTRACT", "SPECIFICATION",
                                                                                  "SUBCONTRACT"):
                body = g["text"].split("\n", 1)[1] if "\n" in g["text"] else g["text"]
                kind = classify_clause(g["heading"], body)
                db.add(ContractClause(document_revision_id=revision.id, main_contract_id=doc.main_contract_id,
                                      clause_reference=g["section"], heading=g["heading"], text=body.strip(),
                                      clause_type=kind, source_passage_id=p.id))
                n_clauses += 1
        job.status, job.completed_at = "COMPLETED", now()
        job.stats = {"passages": len(groups), "clauses": n_clauses, "assertions": n_assert, "flagged": n_flags}
        if revision.business_status != "APPROVED":
            revision.ingestion_state = "REVIEW_REQUIRED" if (n_assert or n_flags or n_clauses) else "VERIFIED"
        else:
            revision.ingestion_state = "PUBLISHED"
        db.commit()
    except Exception as exc:  # leave the job resumable, never half-published
        db.rollback()
        job = db.get(ProcessingJob, job.id)
        job.status, job.error_code = "FAILED", str(exc)[:200]
        db.get(DocumentRevision, revision.id).ingestion_state = "PROCESSING_FAILED"
        db.commit()
    return job


def latest_job_passages(db: Session, revision_id: str):
    job = db.scalar(select(ProcessingJob).where(ProcessingJob.revision_id == revision_id,
                                                ProcessingJob.status == "COMPLETED")
                    .order_by(ProcessingJob.completed_at.desc()))
    if not job:
        return []
    return list(db.scalars(select(Passage).where(Passage.job_id == job.id).order_by(Passage.seq)))
