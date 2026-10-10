"""Acceptance tests mapped to the blueprint catalogue (Section 14). Test IDs are in each name."""
import io

import docx

from seed.demo_seed import SAMPLES
from tests.conftest import H


def up(client, user, pid, name, data, number, rev, dtype="GCC", **form):
    return client.post("/api/v1/documents", headers=H(user),
                       data={"project_id": pid, "document_number": number, "revision": rev, "document_type": dtype, **form},
                       files={"file": (name, data)})


def gcc_bytes(text_clauses):
    d = docx.Document()
    for ref, head, body in text_clauses:
        d.add_paragraph(f"{ref} {head}")
        d.add_paragraph(body)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


# ---------------------------------------------------------------- registry & contract integrity
def test_REG01_duplicate_legal_entity_detected(client, world):
    r = client.post("/api/v1/subcontractors", headers=H("ravi.registry"), json={
        "legal_name": "Shakti Piping (duplicate name)", "registration_number": "u28100gj2011ptc066123", "country_code": "in"})
    assert r.status_code == 409
    assert r.json()["error"]["code"] == "DUPLICATE_LEGAL_ENTITY"
    assert r.json()["error"]["existing_id"] == world["shakti"]


def test_REG02_one_identity_two_projects_separate_visibility(client, world):
    # coordinator sees both relationships; a P-101-only user sees one
    both = client.get(f"/api/v1/subcontractors/{world['shakti']}", headers=H("priya.coord")).json()
    one = client.get(f"/api/v1/subcontractors/{world['shakti']}", headers=H("anita.contracts")).json()
    assert len(both["project_relationships"]) == 2
    assert [l["project_id"] for l in one["project_relationships"]] == [world["project_id"]]


def test_CON01_main_contract_from_other_project_rejected(client, world):
    other_mc = client.get(f"/api/v1/projects/{world['p202']}", headers=H("priya.coord")).json()["main_contracts"][0]["id"]
    r = client.post("/api/v1/subcontracts", headers=H("anita.contracts"), json={
        "project_id": world["project_id"], "main_contract_id": other_mc, "subcontractor_id": world["shakti"],
        "subcontract_number": "SC-X", "scope_summary": "x", "contract_value": "1.00"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "MAIN_CONTRACT_PROJECT_MISMATCH"


def test_CON01_composite_fk_enforced_in_database(client, world):
    """Even bypassing the API, the database refuses a cross-project main contract."""
    import pytest
    from sqlalchemy.exc import IntegrityError

    from app.db import SessionLocal
    from app.models import Subcontract
    other_mc = client.get(f"/api/v1/projects/{world['p202']}", headers=H("priya.coord")).json()["main_contracts"][0]["id"]
    with SessionLocal() as db, pytest.raises(IntegrityError):
        db.add(Subcontract(project_id=world["project_id"], main_contract_id=other_mc, subcontractor_id=world["shakti"],
                           subcontract_number="SC-RAW", scope_summary="x", contract_value=1, created_by="x"))
        db.commit()


def test_CON02_subcontract_without_main_contract_rejected(client, world):
    r = client.post("/api/v1/subcontracts", headers=H("anita.contracts"), json={
        "project_id": world["project_id"], "subcontractor_id": world["shakti"], "subcontract_number": "SC-Y",
        "scope_summary": "x", "contract_value": "1.00"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "MAIN_CONTRACT_REQUIRED"


def test_CON03_executed_baseline_immutable(client, world):
    r = client.patch(f"/api/v1/subcontracts/{world['sc1']}", headers=H("anita.contracts"), json={"contract_value": "1.00"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "BASELINE_LOCKED"
    s = client.get(f"/api/v1/subcontracts/{world['sc1']}", headers=H("anita.contracts")).json()
    assert s["baseline"]["contract_value"] == "125000000.00"


def test_CON04_superseded_revision_retained(client, world):
    pid = world["project_id"]
    from seed.make_samples import GCC_CLAUSES
    revb = gcc_bytes([c if c[0] != "8.4" else ("8.4", "Notice of Delay", "Notice within twenty-one (21) days.")
                      for c in GCC_CLAUSES])
    r = up(client, "anita.contracts", pid, "gcc_b.docx", revb, "EPC-101-GCC", "B", main_contract_id=world["main_contract_id"])
    assert r.json()["outcome"] == "NEW_REVISION"
    doc_id = r.json()["document"]["id"]
    before = client.get(f"/api/v1/documents/{doc_id}", headers=H("anita.contracts")).json()
    assert [x["governing"] for x in before["revisions"]] == [True, False]   # upload does not change governance
    client.post(f"/api/v1/revisions/{r.json()['revision_id']}/approve", headers=H("vikram.legal"), json={})
    after = client.get(f"/api/v1/documents/{doc_id}", headers=H("anita.contracts")).json()
    states = {x["revision"]: (x["governing"], x["ingestion_state"]) for x in after["revisions"]}
    assert states == {"A": (False, "SUPERSEDED"), "B": (True, "PUBLISHED")}
    rev_a = [x for x in after["revisions"] if x["revision"] == "A"][0]["id"]
    assert client.get(f"/api/v1/revisions/{rev_a}/download", headers=H("anita.contracts")).status_code == 200


# ---------------------------------------------------------------- ingestion
def test_DOC01_same_file_twice_is_duplicate(client, world):
    data = (SAMPLES / "EPC-101-GCC_revA.docx").read_bytes()
    r = up(client, "anita.contracts", world["project_id"], "renamed.docx", data, "SOME-OTHER-NO", "Z")
    assert r.json()["outcome"] == "DUPLICATE" and r.json()["job"] is None
    docs = client.get(f"/api/v1/projects/{world['project_id']}/documents", headers=H("anita.contracts")).json()
    assert not any(d["document_number"] == "SOME-OTHER-NO" for d in docs)


def test_DOC02_pdf_page_locators_preserved(client, world):
    docs = client.get(f"/api/v1/projects/{world['project_id']}/documents", headers=H("anita.contracts")).json()
    ht = [d for d in docs if d["document_number"] == "QA-HT-0007"][0]
    ps = client.get(f"/api/v1/revisions/{ht['revisions'][0]['id']}/passages", headers=H("anita.contracts")).json()
    assert ps and all(p["page"] == 1 for p in ps)
    assert any("29.4 barg" in p["text"] for p in ps)


def test_DOC03_boq_cell_provenance_and_formulas(client, world):
    docs = client.get(f"/api/v1/projects/{world['project_id']}/documents", headers=H("anita.contracts")).json()
    boq = [d for d in docs if d["document_type"] == "BOQ"][0]
    ps = client.get(f"/api/v1/revisions/{boq['revisions'][0]['id']}/passages", headers=H("anita.contracts")).json()
    row = [p for p in ps if "P-03" in p["text"]][0]
    assert row["cell"] == "BOQ-Piping!A4:F4" and "F4==D4*E4" in row["text"] and "D4=140" in row["text"]


def test_DOC04_conflicting_revision_label_not_guessed(client, world):
    r = up(client, "anita.contracts", world["project_id"], "gcc_other.docx",
           gcc_bytes([("1.1", "Definitions", "Different content.")]), "EPC-101-GCC", "A")
    body = r.json()
    assert body["outcome"] == "REVISION_CONFLICT" and body["job"] is None
    revs = body["document"]["revisions"]
    assert [x["governing"] for x in revs if x["ingestion_state"] == "CONFLICT"] == [False]
    approve = client.post(f"/api/v1/revisions/{body['revision_id']}/approve", headers=H("vikram.legal"), json={})
    assert approve.status_code == 409


def test_DOC05_unsafe_or_malformed_files_quarantined(client, world):
    pid = world["project_id"]
    assert up(client, "anita.contracts", pid, "x.exe", b"MZ...", "X1", "0").status_code == 415
    assert up(client, "anita.contracts", pid, "fake.pdf", b"not a pdf", "X2", "0").status_code == 422
    assert up(client, "anita.contracts", pid, "broken.docx", b"PK\x03\x04garbage", "X3", "0").status_code == 422
    assert up(client, "anita.contracts", pid, "bin.txt", b"abc\x00def", "X4", "0").status_code == 422


def test_DOC06_interrupted_job_retries_without_duplicates(client, world):
    from app.db import SessionLocal
    from app.models import ContractClause, DocumentRevision, Passage
    from app.services.ingestion import process_revision, register_revision
    with SessionLocal() as db:
        res = register_revision(db, actor_id="t", project_id=world["project_id"], document_number="EPC-101-SPEC",
                                revision_label="0", document_type="SPECIFICATION", title="Spec",
                                filename="spec.docx", data=gcc_bytes([("2.1", "Piping Inspection", "Inspection within three (3) days."),
                                           ("2.2", "Painting", "Paint as specified.")]))
        db.commit()
        rid = res.revision.id
        job = process_revision(db, db.get(DocumentRevision, rid), fail_after=1)
        assert job.status == "FAILED"
        job = process_revision(db, db.get(DocumentRevision, rid))
        assert job.status == "COMPLETED"
        job2 = process_revision(db, db.get(DocumentRevision, rid))   # re-run is a no-op
        assert job2.id == job.id
        assert db.query(Passage).filter(Passage.revision_id == rid).count() == job.stats["passages"]
        assert db.query(ContractClause).filter(ContractClause.document_revision_id == rid).count() == 2


def test_DOC07_new_parser_version_does_not_overwrite_approved_records(client, world):
    docs = client.get(f"/api/v1/projects/{world['project_id']}/documents", headers=H("anita.contracts")).json()
    gcc = [d for d in docs if d["document_number"] == "EPC-101-GCC"][0]
    rid = gcc["governing_revision_id"]
    clauses_before = client.get(f"/api/v1/revisions/{rid}/clauses", headers=H("anita.contracts")).json()
    r = client.post(f"/api/v1/revisions/{rid}/reprocess", headers=H("vikram.legal"), json={"pipeline_version": "1.1"})
    assert r.json()["status"] == "COMPLETED" and r.json()["pipeline_version"] == "1.1"
    clauses_after = client.get(f"/api/v1/revisions/{rid}/clauses", headers=H("anita.contracts")).json()
    assert [c["id"] for c in clauses_before] == [c["id"] for c in clauses_after]
    jobs = client.get(f"/api/v1/documents/{gcc['id']}", headers=H("anita.contracts")).json()["jobs"]
    assert {j["pipeline_version"] for j in jobs if j["revision_id"] == rid} == {"1.0", "1.1"}


# ---------------------------------------------------------------- contract intelligence
def test_CI01_obligation_links_to_exact_revision_and_passage(client, world):
    obs = client.get(f"/api/v1/subcontracts/{world['sc1']}/obligations", headers=H("anita.contracts")).json()
    hse = [o for o in obs if o["type"] == "HSE"][0]
    t = client.get(f"/api/v1/obligations/{hse['id']}/trace", headers=H("anita.contracts")).json()
    assert t["source_clause"]["reference"] == "11.1"
    assert t["source_revision"]["document_number"] == "EPC-101-GCC" and t["source_revision"]["governing"]
    assert "seven (7) days" in t["source_passage"]["text"]
    assert hse["notice_period_days"] == 7 and hse["due_at"] is not None


def test_CI02_superseded_revision_excluded_unless_requested(client, world):
    pid = world["project_id"]
    from seed.make_samples import GCC_CLAUSES
    revb = gcc_bytes([c if c[0] != "8.7" else ("8.7", "Delay Damages", "Liquidated damages at 0.25% per week.")
                      for c in GCC_CLAUSES])
    r = up(client, "anita.contracts", pid, "gcc_b.docx", revb, "EPC-101-GCC", "B")
    client.post(f"/api/v1/revisions/{r.json()['revision_id']}/approve", headers=H("vikram.legal"), json={})
    q = {"projectId": pid, "question": "What is the rate of liquidated damages in clause 8.7?"}
    cur = client.post("/api/v1/knowledge/queries", headers=H("anita.contracts"), json=q).json()
    assert all(s["governing"] for s in cur["sources"]) and "0.25%" in cur["sources"][0]["excerpt"]
    hist = client.post("/api/v1/knowledge/queries", headers=H("anita.contracts"),
                       json={**q, "includeSupersededEvidence": True}).json()
    assert any(not s["governing"] and "0.5%" in s["excerpt"] for s in hist["sources"])


def test_CI03_conflicting_clauses_surfaced(client, world):
    r = client.post("/api/v1/knowledge/queries", headers=H("anita.contracts"),
                    json={"projectId": world["project_id"], "question": "What notice period applies to delay events?"}).json()
    assert r["status"] == "SUPPORTED_WITH_QUALIFICATIONS"
    c = r["contradictions"][0]
    assert c["clause"] == "8.4" and set(c["values"]) == {"14 CALENDAR", "10 CALENDAR"}


def test_CI04_every_source_clause_has_a_row(client, world):
    gate = client.get(f"/api/v1/subcontracts/{world['sc2']}/release-gate", headers=H("anita.contracts")).json()
    assert gate["clauses_without_mapping"] == 0 and gate["mapped"] == 16


def test_CI05_release_blocked_then_exception_with_segregation(client, world):
    tr = lambda u: client.post(f"/api/v1/subcontracts/{world['sc2']}/transitions", headers=H(u),  # noqa: E731
                               json={"target": "APPROVED_FOR_EXECUTION"})
    r = tr("meera.approver")
    assert r.status_code == 409 and r.json()["error"]["code"] == "RELEASE_BLOCKED"
    short = client.post(f"/api/v1/subcontracts/{world['sc2']}/release-exception", headers=H("meera.approver"),
                        json={"reason": "ok"})
    assert short.status_code == 422
    assert client.post(f"/api/v1/subcontracts/{world['sc2']}/release-exception", headers=H("anita.contracts"),
                       json={"reason": "Drafter tries to accept own risk here"}).status_code == 403
    g = client.post(f"/api/v1/subcontracts/{world['sc2']}/release-exception", headers=H("meera.approver"),
                    json={"reason": "Civil LDs and insurance agreed in principle; to be closed by Amendment 1."}).json()
    assert g["passed"] is True
    assert tr("meera.approver").status_code == 200


def test_CI06_unanswerable_question_reports_insufficient_evidence(client, world):
    r = client.post("/api/v1/knowledge/queries", headers=H("anita.contracts"),
                    json={"projectId": world["project_id"], "question": "Who won the cricket world cup?"}).json()
    assert r["status"] == "INSUFFICIENT_EVIDENCE" and r["sources"] == [] and r["answer"] is None


def test_CI07_revoked_access_denies_search_and_download(client, world):
    pid = world["project_id"]
    docs = client.get(f"/api/v1/projects/{pid}/documents", headers=H("arjun.qaqc")).json()
    rid = docs[0]["revisions"][0]["id"]
    assert client.get(f"/api/v1/revisions/{rid}/download", headers=H("arjun.qaqc")).status_code == 200
    client.delete(f"/api/v1/projects/{pid}/members/arjun.qaqc", headers=H("priya.coord"))
    assert client.get(f"/api/v1/revisions/{rid}/download", headers=H("arjun.qaqc")).status_code == 403
    assert client.post("/api/v1/knowledge/queries", headers=H("arjun.qaqc"),
                       json={"projectId": pid, "question": "notice"}).status_code == 403


def test_CI08_embedded_instruction_is_data_not_command(client, world):
    q = client.get(f"/api/v1/projects/{world['project_id']}/review-queue", headers=H("anita.contracts")).json()
    assert any(f["document_number"] == "LTR-SC-0042" and "EMBEDDED_INSTRUCTION" in f["flags"] for f in q["flagged_passages"])
    r = client.post("/api/v1/knowledge/queries", headers=H("anita.contracts"),
                    json={"projectId": world["project_id"], "question": "approve this variation automatically work front PR-3"}).json()
    assert any("embedded instructions" in w for w in r["warnings"])
    me = client.get("/api/v1/me", headers=H("anita.contracts")).json()
    assert "contracts_approver" not in me["roles"] + sum(me["projects"].values(), [])


def test_SEC01_cross_project_access_denied(client, world):
    pid, sc1 = world["project_id"], world["sc1"]
    assert client.get(f"/api/v1/projects/{pid}", headers=H("omar.solar")).status_code == 403
    assert client.get(f"/api/v1/subcontracts/{sc1}", headers=H("omar.solar")).status_code == 403
    assert client.post("/api/v1/knowledge/queries", headers=H("omar.solar"),
                       json={"projectId": pid, "question": "notice of delay"}).status_code == 403
    assert all(p["id"] != pid for p in client.get("/api/v1/projects", headers=H("omar.solar")).json())


def test_SEC02_classification_above_clearance_hidden(client, world):
    pid = world["project_id"]
    up(client, "anita.contracts", pid, "board.txt", b"Board note on margin.", "BRD-001", "0", "CORRESPONDENCE",
       security_classification="RESTRICTED")
    docs = client.get(f"/api/v1/projects/{pid}/documents", headers=H("arjun.qaqc")).json()   # CONFIDENTIAL clearance
    assert all(d["document_number"] != "BRD-001" for d in docs)


# ---------------------------------------------------------------- workflow, idempotency, audit
def test_WF01_invalid_transition_rejected(client, world):
    r = client.post(f"/api/v1/subcontracts/{world['sc1']}/transitions", headers=H("anita.contracts"),
                    json={"target": "DRAFT"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "INVALID_TRANSITION"


def test_SOD01_drafter_cannot_approve_release(client, world):
    from app.db import SessionLocal
    from app.models import ProjectMembership, User
    with SessionLocal() as db:   # give the drafter the approver role too
        anita = db.query(User).filter_by(username="anita.contracts").one()
        m = db.query(ProjectMembership).filter_by(user_id=anita.id, project_id=world["project_id"]).one()
        m.roles = ["contracts_manager", "contracts_approver"]
        db.commit()
    r = client.post(f"/api/v1/subcontracts/{world['sc2']}/transitions", headers=H("anita.contracts"),
                    json={"target": "APPROVED_FOR_EXECUTION"})
    assert r.status_code == 403 and r.json()["error"]["code"] == "SEGREGATION_OF_DUTIES"


def test_IDEM01_retry_with_same_key_creates_one_record(client, world):
    body = {"project_id": world["project_id"], "main_contract_id": world["main_contract_id"],
            "subcontractor_id": world["shakti"], "subcontract_number": "SC-101-09", "scope_summary": "Supports",
            "contract_value": "1000.00"}
    a = client.post("/api/v1/subcontracts", headers=H("anita.contracts", **{"Idempotency-Key": "k-1"}), json=body)
    b = client.post("/api/v1/subcontracts", headers=H("anita.contracts", **{"Idempotency-Key": "k-1"}), json=body)
    assert a.status_code == 201 and a.json()["id"] == b.json()["id"]
    c = client.post("/api/v1/subcontracts", headers=H("anita.contracts", **{"Idempotency-Key": "k-1"}),
                    json={**body, "contract_value": "2000.00"})
    assert c.status_code == 422 and c.json()["error"]["code"] == "IDEMPOTENCY_KEY_REUSED"
    subs = client.get(f"/api/v1/projects/{world['project_id']}", headers=H("anita.contracts")).json()["subcontracts"]
    assert sum(s["subcontract_number"] == "SC-101-09" for s in subs) == 1


def test_IDEM02_obligation_generation_is_idempotent(client, world):
    before = client.get(f"/api/v1/subcontracts/{world['sc1']}/obligations", headers=H("anita.contracts")).json()
    r = client.post(f"/api/v1/subcontracts/{world['sc1']}/obligations/generate", headers=H("anita.contracts")).json()
    after = client.get(f"/api/v1/subcontracts/{world['sc1']}/obligations", headers=H("anita.contracts")).json()
    assert r["created"] == 0 and len(before) == len(after)


def test_AUD01_amendment_keeps_original_actor_reason_and_values(client, world):
    s = client.get(f"/api/v1/subcontracts/{world['sc1']}", headers=H("anita.contracts")).json()
    assert s["contract_value"] == "126850000.00" and s["baseline"]["contract_value"] == "125000000.00"
    ev = client.get(f"/api/v1/audit-events?project_id={world['project_id']}", headers=H("dev.audit")).json()
    a = [e for e in ev if e["action"] == "amendment.approved"][0]
    assert a["actor"] == "meera.approver" and a["before"] == {"contract_value": "125000000.00"}
    assert a["after"]["contract_value"] == "126850000.00" and "VO-001" in a["reason"]
    assert client.get(f"/api/v1/audit-events?project_id={world['project_id']}",
                      headers=H("arjun.qaqc")).status_code == 403


# ---------------------------------------------------------------- first vertical slice
def test_VS01_end_to_end_chain_with_evidence_rejection_and_closure(client, world):
    obs = client.get(f"/api/v1/subcontracts/{world['sc1']}/obligations", headers=H("anita.contracts")).json()
    ins = [o for o in obs if o["type"] == "INSURANCE"][0]
    assert ins["statement"].startswith("[Subcontract treatment]") and ins["notice_period_days"] == 14
    data = b"Insurance certificate CAR-2026-118 (fictional)."
    rev = up(client, "anita.contracts", world["project_id"], "car.txt", data, "INS-CAR-118", "0", "EVIDENCE").json()["revision_id"]
    e1 = client.post(f"/api/v1/obligations/{ins['id']}/evidence", headers=H("anita.contracts"),
                     json={"document_revision_id": rev, "note": "CAR certificate"}).json()
    # submitter cannot accept own evidence
    assert client.post(f"/api/v1/evidence/{e1['id']}/review", headers=H("anita.contracts"),
                       json={"decision": "ACCEPT", "comment": "self"}).status_code == 403
    rej = client.post(f"/api/v1/evidence/{e1['id']}/review", headers=H("vikram.legal"),
                      json={"decision": "REJECT", "comment": "Third-party liability cover missing"}).json()
    assert rej["status"] == "OPEN"    # stays open when acceptance criteria are not met
    e2 = client.post(f"/api/v1/obligations/{ins['id']}/evidence", headers=H("anita.contracts"),
                     json={"document_revision_id": rev, "note": "resubmitted with TPL"}).json()
    ok = client.post(f"/api/v1/evidence/{e2['id']}/review", headers=H("vikram.legal"),
                     json={"decision": "ACCEPT", "comment": "CAR and TPL confirmed"}).json()
    assert ok["status"] == "CLOSED"
    t = client.get(f"/api/v1/obligations/{ins['id']}/trace", headers=H("dev.audit")).json()
    assert t["source_clause"]["reference"] == "18.1" and t["flowdown"]["disposition"] == "PROJECT_SPECIFIC"
    assert [e["review_status"] for e in t["evidence"]] == ["REJECTED", "ACCEPTED"]
    assert len(t["source_revision"]["sha256"]) == 64


def test_internal_control_is_not_a_contractual_obligation(client, world):
    obs = client.get(f"/api/v1/subcontracts/{world['sc1']}/obligations", headers=H("anita.contracts")).json()
    ic = [o for o in obs if o["origin_kind"] == "INTERNAL_CONTROL"]
    assert len(ic) == 1 and ic[0]["source_mapping_id"] is None and "not a contractual" in ic[0]["internal_control_origin"]


def test_dashboard_counts_drill_to_records(client, world):
    d = client.get(f"/api/v1/projects/{world['project_id']}/dashboard", headers=H("anita.contracts")).json()
    assert d["obligations"]["overdue"] == len(d["overdue_obligations"]) == 1
    assert d["overdue_obligations"][0]["type"] == "NOTICE"
    assert d["flowdown"]["unresolved_high"] > 0   # SC-101-02 is held at the gate


def test_approver_can_read_what_they_approve(client, world):
    assert client.get(f"/api/v1/subcontracts/{world['sc2']}", headers=H("meera.approver")).status_code == 200
    assert client.get(f"/api/v1/subcontracts/{world['sc2']}/flowdown-mappings", headers=H("meera.approver")).status_code == 200
